"""Leitura limitada de ODS/XLSX; fórmulas nunca são executadas."""
from collections import Counter
from io import BytesIO
from pathlib import Path
import unicodedata
from zipfile import ZipFile, BadZipFile
from defusedxml.ElementTree import fromstring
from openpyxl import load_workbook
from pydantic import ValidationError
from .models import Client

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5001
MAX_COLS = 40
NS = {'t':'urn:oasis:names:tc:opendocument:xmlns:table:1.0', 'x':'urn:oasis:names:tc:opendocument:xmlns:text:1.0'}
ALIASES = {'NOME':'nome','ENDERECO':'endereco','NUMERO':'numero','COMPL':'complemento',
 'COMPLEMENTO':'complemento','BAIRRO':'bairro','CIDADE':'cidade','ESTADO':'estado','UF':'estado',
 'CEP':'cep','TELEFONE':'telefone','CGC':'cnpj','CNPJ':'cnpj','EMAIL':'email','E-MAIL':'email',
 'NUMERO_CONTRATO':'numero_contrato','CONTRATO_STATUS':'contrato_status'}


def label(value):
    return ''.join(c for c in unicodedata.normalize('NFD',str(value or '').strip().upper())
                   if unicodedata.category(c) != 'Mn').replace(' ','_')


def read_rows(data: bytes, filename: str):
    if len(data) > MAX_BYTES:
        raise ValueError('Arquivo excede 5 MB.')
    suffix = Path(filename).suffix.lower()
    if suffix not in ('.xlsx','.ods'):
        raise ValueError('Use uma planilha .xlsx ou .ods.')
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 30*1024*1024 or len(archive.infolist()) > 1000:
                raise ValueError('Conteúdo descompactado excede os limites permitidos.')
            if suffix == '.ods':
                root = fromstring(archive.read('content.xml'))
                tables = root.findall('.//t:table',NS)
                if len(tables) != 1:
                    raise ValueError('Envie uma planilha com uma única aba.')
                rows = []
                logical = 0
                for row in tables[0].findall('t:table-row',NS):
                    values = []
                    for cell in row:
                        if cell.tag not in ('{'+NS['t']+'}table-cell','{'+NS['t']+'}covered-table-cell'):
                            continue
                        val = '\n'.join(''.join(p.itertext()) for p in cell.findall('x:p',NS))
                        if cell.get('{'+NS['t']+'}formula') is not None:
                            val = '=FORMULA_NAO_PERMITIDA'
                        repeated = int(cell.get('{'+NS['t']+'}number-columns-repeated','1'))
                        if repeated < 1:
                            raise ValueError('Repetição inválida na planilha.')
                        values.extend([val]*min(repeated,MAX_COLS+1))
                    while values and values[-1] == '':
                        values.pop()
                    repeats = int(row.get('{'+NS['t']+'}number-rows-repeated','1'))
                    if repeats < 1:
                        raise ValueError('Repetição inválida na planilha.')
                    if values:
                        if logical+repeats > MAX_ROWS or len(values)>MAX_COLS:
                            raise ValueError('Limite de 5.000 clientes e 40 colunas excedido.')
                        for i in range(repeats):
                            rows.append((logical+i+1,values))
                    logical += repeats
                return rows
        workbook = load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
        try:
            if len(workbook.worksheets) != 1:
                raise ValueError('Envie uma planilha com uma única aba.')
            sheet = workbook.worksheets[0]
            if (sheet.max_row or 0)>MAX_ROWS or (sheet.max_column or 0)>MAX_COLS:
                raise ValueError('Limite de 5.000 clientes e 40 colunas excedido.')
            return [(i,list(row)) for i,row in enumerate(sheet.iter_rows(values_only=True),1) if any(v is not None for v in row)]
        finally:
            workbook.close()
    except (BadZipFile,KeyError) as exc:
        raise ValueError('Arquivo de planilha inválido.') from exc


def preview(data,filename):
    rows = read_rows(data,filename)
    if not rows:
        raise ValueError('A planilha está vazia.')
    headers = [ALIASES.get(label(c)) for c in rows[0][1]]
    missing = {'nome','cnpj','endereco','cidade','estado'} - set(headers)
    if missing:
        raise ValueError('Colunas obrigatórias ausentes: '+', '.join(sorted(missing)))
    mapped = [x for x in headers if x]
    if len(mapped) != len(set(mapped)):
        raise ValueError('Colunas repetidas para o mesmo campo.')
    result = []
    for number,values in rows[1:]:
        raw = {}
        errors = []
        for i,field in enumerate(headers):
            if not field:
                continue
            value = values[i] if i < len(values) else None
            if value is None or str(value).strip() == '':
                continue
            if isinstance(value,str) and value.startswith('='):
                errors.append(f'{field}: substitua a fórmula por um valor.')
                continue
            if isinstance(value,(int,float)) and field in ('cnpj','cep'):
                errors.append(f'{field}: formate como texto para preservar zeros iniciais.')
                continue
            raw[field] = str(int(value)) if isinstance(value,float) and value.is_integer() else str(value).strip()
        if raw.get('numero_contrato') and not raw.get('contrato_status'):
            raw['contrato_status'] = 'com_contrato'
        normalized = None
        try:
            normalized = Client.model_validate(raw).model_dump(mode='json')
        except ValidationError as exc:
            errors.extend(f'{".".join(map(str,e["loc"])) or "cliente"}: {e["msg"]}' for e in exc.errors())
        result.append({'linha':number,'dados':normalized,'erros':errors})
    counts = Counter(r['dados']['cnpj'] for r in result if r['dados'])
    for row in result:
        if row['dados'] and counts[row['dados']['cnpj']]>1:
            row['erros'].append('CNPJ repetido na planilha; mantenha uma única linha por cliente.')
    return result
